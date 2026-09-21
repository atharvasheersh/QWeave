"""Build the editable Word manuscript from the maintained LaTeX source."""

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "manuscript.tex"
OUTPUT = ROOT / "docx" / "QWeave_Manuscript.docx"
LEARNING_CURVE = ROOT.parent / "results" / "benchmark_v2" / "figures" / "learning_curves.png"

CITATIONS = {
    "li2019sabre": 1, "pozzi2020rl": 2, "sinha2022gnn": 3,
    "nannicini2022ip": 4, "ibmtranspiler": 5, "ibmsabre": 6,
}
REFERENCES = [
    "G. Li, Y. Ding, and Y. Xie. Tackling the Qubit Mapping Problem for NISQ-Era Quantum Devices. ASPLOS, 2019. arXiv:1809.02573.",
    "M. G. Pozzi, S. J. Herbert, A. Sengupta, and R. D. Mullins. Using Reinforcement Learning to Perform Qubit Routing in Quantum Compilers. arXiv:2007.15957, 2020.",
    "A. Sinha, U. Azad, and H. Singh. Qubit Routing using Graph Neural Network aided Monte Carlo Tree Search. AAAI, 2022. arXiv:2104.01992.",
    "G. Nannicini, L. S. Bishop, O. Gunluk, and P. Jurcevic. Optimal Qubit Assignment and Routing via Integer Programming. ACM Transactions on Quantum Computing, 2022.",
    "IBM Quantum. Transpiler stages. IBM Quantum Documentation, 2026.",
    "IBM Quantum. SabreLayout API reference. IBM Quantum Documentation, 2026.",
]


def _field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, text, end))


def _set_columns(section, count: int) -> None:
    section_properties = section._sectPr
    columns = section_properties.xpath("./w:cols")
    element = columns[0] if columns else OxmlElement("w:cols")
    element.set(qn("w:num"), str(count))
    element.set(qn("w:space"), "500")
    if not columns:
        section_properties.append(element)


def _clean(text: str) -> str:
    def citation(match) -> str:
        numbers = sorted(CITATIONS[key] for key in match.group(1).split(","))
        return "[" + ", ".join(str(number) for number in numbers) + "]"

    replacements = {
        r"Eq.~\eqref{eq:objective}": "Equation (1)",
        r"Table~\ref{tab:test-results}": "Table 1",
        r"Figure~\ref{fig:learning-curves}": "Figure 1",
        r"\eqref{eq:equivalence}": "Equation (3)",
        r"\rightarrow": "->", r"\ldots": "...", r"\emph": "",
        r"\texttt": "", r"\textsc": "", r"\small": "",
        r"\{": "{", r"\}": "}", r"\%": "%", r"\_": "-", r"~": " ",
    }
    text = re.sub(r"\\cite\{([^}]+)\}", citation, text)
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"\\(?:label|ref)\{[^}]+\}", "", text)
    text = re.sub(r"\\[A-Za-z]+", "", text)
    text = (text.replace("{", "").replace("}", "").replace("$", "")
            .replace("--", "-").replace("_", ""))
    return " ".join(text.split())


def _add_body_paragraph(document: Document, text: str) -> None:
    cleaned = _clean(text)
    if not cleaned:
        return
    paragraph = document.add_paragraph(cleaned)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(0.35)
    paragraph.paragraph_format.space_after = Pt(3)


def _add_results_table(document: Document) -> None:
    wide = document.add_section(WD_SECTION.CONTINUOUS)
    wide.page_width, wide.page_height = Cm(21), Cm(29.7)
    wide.top_margin = wide.bottom_margin = Cm(1.7)
    wide.left_margin = wide.right_margin = Cm(1.8)
    _set_columns(wide, 1)
    caption = document.add_paragraph("Table 1  Benchmark-v2 test summary. Runtime is median wall time in seconds.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].bold = True
    table = document.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ("Method", "Depth", "SWAP", "Runtime", "Fallback")
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "1F4E78")
        cell._tc.get_or_add_tcPr().append(shading)
    rows = (
        ("Basic", "48.5", "11.5", "0.0022", "0"),
        ("Weighted", "36.0", "0.0", "0.0036", "0"),
        ("SABRE", "46.0", "7.0", "0.0028", "0"),
        ("GNN-PPO", "37.5", "0.0", "0.3473", "4543"),
        ("No-message PPO", "36.0", "0.0", "0.1040", "742"),
        ("Untrained GNN", "198.0", "193.0", "0.6681", "8059"),
    )
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.paragraphs[0].alignment = (WD_ALIGN_PARAGRAPH.LEFT
                                            if cell is cells[0]
                                            else WD_ALIGN_PARAGRAPH.CENTER)
            if row_index % 2:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "DCE6F1")
                cell._tc.get_or_add_tcPr().append(shading)
    document.add_paragraph()
    columns = document.add_section(WD_SECTION.CONTINUOUS)
    columns.page_width, columns.page_height = Cm(21), Cm(29.7)
    columns.top_margin = columns.bottom_margin = Cm(1.7)
    columns.left_margin = columns.right_margin = Cm(1.8)
    _set_columns(columns, 2)


def _add_learning_figure(document: Document) -> None:
    if not LEARNING_CURVE.exists():
        raise FileNotFoundError(f"missing benchmark figure: {LEARNING_CURVE}")
    wide = document.add_section(WD_SECTION.CONTINUOUS)
    wide.page_width, wide.page_height = Cm(21), Cm(29.7)
    wide.top_margin = wide.bottom_margin = Cm(1.7)
    wide.left_margin = wide.right_margin = Cm(1.8)
    _set_columns(wide, 1)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(LEARNING_CURVE), width=Cm(16.5))
    caption = document.add_paragraph(
        "Figure 1  Five-seed learning curves. Lines show the median recent "
        "episode return and bands show the interquartile range.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    columns = document.add_section(WD_SECTION.CONTINUOUS)
    columns.page_width, columns.page_height = Cm(21), Cm(29.7)
    columns.top_margin = columns.bottom_margin = Cm(1.7)
    columns.left_margin = columns.right_margin = Cm(1.8)
    _set_columns(columns, 2)


def build() -> Path:
    source = SOURCE.read_text(encoding="utf-8")
    title = re.search(r"\\title\{(.+?)\}", source).group(1)
    author = re.search(r"\\author\{(.+?)\}", source).group(1)
    date = re.search(r"\\date\{(.+?)\}", source).group(1)
    abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", source, re.S).group(1)
    body = source.split(r"\end{abstract}", 1)[1].split(r"\bibliographystyle", 1)[0]
    body = re.sub(r"\\begin\{table\}.*?\\end\{table\}", "\n\n[[RESULTS_TABLE]]\n\n", body, flags=re.S)
    body = re.sub(r"\\begin\{figure\*\}.*?\\end\{figure\*\}",
                  "\n\n[[LEARNING_FIGURE]]\n\n", body, flags=re.S)
    equations = [
        "J(m) = sum over (i,j) in E_C of w_ij d_H(m(i),m(j)).   (1)",
        "r_i = d_i + max over successors j of r_j;   s_i = max over predecessors j of (s_j + d_j).   (2)",
        "U_routed P_initial = P_final U_C.   (3)",
    ]
    for equation in equations:
        body = re.sub(r"\\begin\{equation\}.*?\\end\{equation\}",
                      f"\n\n[[EQUATION:{equation}]]\n\n", body, count=1, flags=re.S)

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.7)
    section.left_margin = section.right_margin = Cm(1.8)
    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(9)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for name, size in (("Title", 16), ("Heading 1", 11)):
        styles[name].font.name = "Times New Roman"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    title_properties = styles["Title"]._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Title"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = paragraph.add_run(_clean(title))
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(16)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph = document.add_paragraph(_clean(author))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.runs[0].bold = True
    paragraph = document.add_paragraph(_clean(date))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading = document.add_paragraph("Abstract")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.runs[0].bold = True
    _add_body_paragraph(document, abstract)

    columns = document.add_section(WD_SECTION.CONTINUOUS)
    columns.page_width, columns.page_height = Cm(21), Cm(29.7)
    columns.top_margin = columns.bottom_margin = Cm(1.7)
    columns.left_margin = columns.right_margin = Cm(1.8)
    _set_columns(columns, 2)

    tokens = re.split(r"(\\section\{[^}]+\}|\[\[RESULTS_TABLE\]\]|\[\[LEARNING_FIGURE\]\]|\[\[EQUATION:.*?\]\])", body)
    for token in tokens:
        token = token.strip()
        if not token:
            continue
        section_match = re.fullmatch(r"\\section\{([^}]+)\}", token)
        if section_match:
            document.add_heading(_clean(section_match.group(1)), level=1)
        elif token == "[[RESULTS_TABLE]]":
            _add_results_table(document)
        elif token == "[[LEARNING_FIGURE]]":
            _add_learning_figure(document)
        elif token.startswith("[[EQUATION:"):
            equation = token[len("[[EQUATION:"):-2]
            paragraph = document.add_paragraph(equation)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.runs[0].italic = True
        else:
            for block in re.split(r"\n\s*\n", token):
                _add_body_paragraph(document, block)

    document.add_heading("References", level=1)
    for number, reference in enumerate(REFERENCES, start=1):
        paragraph = document.add_paragraph(f"[{number}] {reference}")
        paragraph.paragraph_format.left_indent = Cm(0.35)
        paragraph.paragraph_format.first_line_indent = Cm(-0.35)
        paragraph.paragraph_format.space_after = Pt(2)

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(footer, "PAGE")
    document.core_properties.title = _clean(title)
    document.core_properties.author = "QWeave Research Team"
    document.core_properties.subject = "Correctness-first qubit mapping and routing study"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
